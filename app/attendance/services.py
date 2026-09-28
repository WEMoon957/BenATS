"""考勤业务逻辑：打卡解析、员工匹配、导入处理、跨日疑似检测与结果重算。

移植自西鸣人事管理系统（Django），改为在 SQLite 数据层上执行，保持核算口径一致。
"""

from __future__ import annotations

import calendar
import hashlib
import re
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from .db import AttendanceStore, _now

TIME_TOKEN = re.compile(r"^(?P<next>次日)?(?P<hour>\d{1,2}):(?P<minute>\d{2})$")

STATUS_PENDING = "pending"
STATUS_PROCESSING = "processing"
STATUS_COMPLETED = "completed"
STATUS_FAILED = "failed"

MODE_STANDARD = "standard"
MODE_FLEXIBLE = "flexible"
MODE_EXEMPT = "exempt"


def normalize_name(value) -> str:
    return re.sub(r"\s+|[（(]已离职[）)]", "", str(value or "")).strip()


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def file_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def parse_punches(value) -> list[dict]:
    raw = cell_text(value)
    if not raw or raw == "-":
        return []
    punches: list[dict] = []
    for token in re.split(r"[\r\n]+", raw):
        token = token.strip()
        if not token or token == "-":
            continue
        match = TIME_TOKEN.match(token)
        if not match:
            punches.append({"text": token, "minutes": None, "next_day": False, "valid": False})
            continue
        minutes = int(match.group("hour")) * 60 + int(match.group("minute"))
        next_day = bool(match.group("next"))
        punches.append(
            {
                "text": token,
                "minutes": minutes + (1440 if next_day else 0),
                "clock_minutes": minutes,
                "next_day": next_day,
                "valid": True,
            }
        )
    return punches


def _employee_maps(store: AttendanceStore):
    employees = store.query(
        "SELECT e.*, p.mode AS policy_mode, p.cross_day_cutoff_minutes AS policy_cutoff "
        "FROM employee e LEFT JOIN attendance_policy p ON p.id = e.attendance_policy_id "
        "WHERE e.active = 1 ORDER BY e.department, e.employee_no",
        table="employee",
    )
    by_no: dict[str, dict] = {}
    by_name: dict[str, dict] = {}
    by_alias: dict[str, dict] = {}
    for employee in employees:
        if employee.get("employee_no", "").strip():
            by_no[employee["employee_no"].strip()] = employee
        by_name.setdefault(normalize_name(employee.get("name")), employee)
        for alias in employee.get("aliases") or []:
            by_alias.setdefault(normalize_name(alias), employee)
    return employees, by_no, by_name, by_alias


def match_employee(employee_no, source_name, by_no, by_name, by_alias):
    if employee_no and employee_no in by_no:
        return by_no[employee_no], "employee_no"
    normalized = normalize_name(source_name)
    if normalized in by_name:
        return by_name[normalized], "name"
    if normalized in by_alias:
        return by_alias[normalized], "alias"
    return None, "unmatched"


def _locate_punch_sheet(workbook):
    if "飞书打卡" in workbook.sheetnames:
        return workbook["飞书打卡"]
    for sheet in workbook.worksheets:
        headers = [cell_text(cell.value) for cell in sheet[1][:4]]
        if headers[:2] == ["姓名", "组织名称"] and "工号" in headers:
            return sheet
    raise ValueError("没有找到打卡明细页：需要包含“姓名、组织名称、工号、考勤规则”表头")


def _day_columns(sheet, year, month):
    max_day = calendar.monthrange(year, month)[1]
    columns: dict[int, int] = {}
    for cell in sheet[1]:
        header = cell_text(cell.value)
        match = re.match(r"^(\d{1,2})", header)
        if not match:
            continue
        day = int(match.group(1))
        if 1 <= day <= max_day:
            columns[day] = cell.column
    if not columns:
        raise ValueError("打卡明细页没有识别到日期列")
    return columns


def _cutoff_for(employee: dict | None, policy_cutoff=None) -> int:
    if employee and employee.get("policy_cutoff") is not None:
        return int(employee["policy_cutoff"])
    if policy_cutoff is not None:
        return int(policy_cutoff)
    return 180


def process_import_batch(store: AttendanceStore, batch_id: int) -> dict:
    batch = store.query_one("SELECT * FROM import_batch WHERE id = ?", (batch_id,), table="import_batch")
    if not batch:
        raise ValueError("导入批次不存在")

    store.execute(
        "UPDATE import_batch SET status = ?, error_message = '', completed_at = NULL WHERE id = ?",
        (STATUS_PROCESSING, batch_id),
    )
    try:
        workbook = load_workbook(batch["source_file_path"], data_only=True, read_only=False)
        sheet = _locate_punch_sheet(workbook)
        day_columns = _day_columns(sheet, batch["year"], batch["month"])
        _employees, by_no, by_name, by_alias = _employee_maps(store)

        with store.transaction() as tx:
            tx.execute("DELETE FROM raw_punch_day WHERE batch_id = ?", (batch_id,))
            total_rows = 0
            matched_rows = 0
            unmatched_rows = 0
            days_for_rows: list[dict[int, dict]] = []

            for row_number in range(2, sheet.max_row + 1):
                source_name = cell_text(sheet.cell(row_number, 1).value)
                if not source_name:
                    continue
                total_rows += 1
                employee_no = cell_text(sheet.cell(row_number, 3).value)
                employee, match_status = match_employee(
                    employee_no, source_name, by_no, by_name, by_alias
                )
                if employee:
                    matched_rows += 1
                else:
                    unmatched_rows += 1
                organization = cell_text(sheet.cell(row_number, 2).value)
                attendance_rule = cell_text(sheet.cell(row_number, 4).value)
                days_for_row: dict[int, dict] = {}
                for day, column in day_columns.items():
                    work_date = date(batch["year"], batch["month"], day)
                    raw_value = cell_text(sheet.cell(row_number, column).value)
                    punches = parse_punches(raw_value)
                    raw_id = tx.execute(
                        "INSERT INTO raw_punch_day (batch_id, employee_id, source_row, employee_no, "
                        "source_name, organization, attendance_rule, work_date, raw_value, punches, "
                        "has_punch, effective_has_punch, match_status, is_cross_day_suspicion) "
                        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)",
                        (
                            batch_id,
                            employee["id"] if employee else None,
                            row_number,
                            employee_no,
                            source_name,
                            organization,
                            attendance_rule,
                            work_date.isoformat(),
                            raw_value,
                            json_dumps(punches),
                            1 if punches else 0,
                            1 if punches else 0,
                            match_status,
                        ),
                    )
                    raw_day = {
                        "id": raw_id,
                        "employee_id": employee["id"] if employee else None,
                        "employee": employee,
                        "work_date": work_date,
                        "punches": punches,
                        "source_row": row_number,
                        "source_name": source_name,
                        "employee_no": employee_no,
                        "organization": organization,
                        "attendance_rule": attendance_rule,
                    }
                    days_for_row[day] = raw_day
                days_for_rows.append(days_for_row)

            # 跨日疑似检测
            suspicion_count = 0
            for days_for_row in days_for_rows:
                for day in sorted(days_for_row):
                    if day <= 1:
                        continue
                    current = days_for_row[day]
                    previous = days_for_row.get(day - 1)
                    if not previous or len(current["punches"]) != 1 or not previous["punches"]:
                        continue
                    punch = current["punches"][0]
                    if not punch.get("valid") or punch.get("next_day"):
                        continue
                    cutoff = _cutoff_for(current.get("employee"))
                    if punch.get("clock_minutes", 10_000) > cutoff:
                        continue
                    previous_has_late_punch = any(
                        item.get("valid")
                        and not item.get("next_day")
                        and item.get("clock_minutes", 0) >= 12 * 60
                        for item in previous["punches"]
                    )
                    if not previous_has_late_punch:
                        continue
                    tx.execute(
                        "UPDATE raw_punch_day SET is_cross_day_suspicion = 1, effective_has_punch = 0 WHERE id = ?",
                        (current["id"],),
                    )
                    tx.execute(
                        "INSERT INTO cross_day_suspicion (batch_id, raw_day_id, employee_id, previous_date, "
                        "work_date, punch_text, reason, status, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        (
                            batch_id,
                            current["id"],
                            current.get("employee_id"),
                            previous["work_date"].isoformat(),
                            current["work_date"].isoformat(),
                            punch["text"],
                            f"当天仅一条凌晨打卡，且前一天存在下午或晚间打卡（截止 {cutoff // 60:02d}:{cutoff % 60:02d}）",
                            STATUS_PENDING,
                            _now(),
                        ),
                    )
                    suspicion_count += 1

        recalculate_batch(store, batch_id)

        store.execute(
            "UPDATE import_batch SET status = ?, total_rows = ?, matched_rows = ?, unmatched_rows = ?, "
            "suspicion_count = ?, completed_at = ? WHERE id = ?",
            (STATUS_COMPLETED, total_rows, matched_rows, unmatched_rows, suspicion_count, _now(), batch_id),
        )
    except Exception as exc:
        store.execute(
            "UPDATE import_batch SET status = ?, error_message = ? WHERE id = ?",
            (STATUS_FAILED, str(exc), batch_id),
        )
        raise
    return store.query_one("SELECT * FROM import_batch WHERE id = ?", (batch_id,), table="import_batch")


def json_dumps(value) -> str:
    import json

    return json.dumps(value, ensure_ascii=False)


def _r(value) -> float:
    return round(float(value or 0), 2)


def recalculate_result(store: AttendanceStore, result_id: int) -> dict:
    result = store.query_one(
        "SELECT r.*, e.attendance_policy_id, e.expected_days_override, p.mode AS policy_mode "
        "FROM attendance_result r "
        "JOIN employee e ON e.id = r.employee_id "
        "LEFT JOIN attendance_policy p ON p.id = e.attendance_policy_id "
        "WHERE r.id = ?",
        (result_id,),
        table="attendance_result",
    )
    if not result:
        raise ValueError("考勤结果不存在")
    batch = store.query_one("SELECT * FROM import_batch WHERE id = ?", (result["batch_id"],))
    punch_days = store.query(
        "SELECT COUNT(*) AS c FROM raw_punch_day WHERE batch_id = ? AND employee_id = ? AND effective_has_punch = 1",
        (result["batch_id"], result["employee_id"]),
    )[0]["c"]

    mode = result.get("policy_mode") or MODE_STANDARD
    due_days = _r(result.get("expected_days_override") or batch["default_expected_days"])
    days_in_month = calendar.monthrange(batch["year"], batch["month"])[1]
    rest_days = max(0.0, round(float(days_in_month) - due_days, 2))

    if mode in {MODE_EXEMPT, MODE_FLEXIBLE}:
        base_actual = due_days
        base_rule = "免考勤/弹性工作：按应出勤天数正常计薪"
    else:
        base_actual = float(punch_days)
        base_rule = "标准规则：空白或“-”为休息，有打卡为出勤；疑似跨日未确认前不计入当天"

    adjustment = _r(result.get("adjustment_days"))
    actual_days = round(base_actual + adjustment, 2)
    pending_suspicions = store.query(
        "SELECT COUNT(*) AS c FROM cross_day_suspicion WHERE batch_id = ? AND employee_id = ? AND status = 'pending'",
        (result["batch_id"], result["employee_id"]),
    )[0]["c"]

    status = result["status"]
    if status != "approved":
        if pending_suspicions or (mode == MODE_STANDARD and actual_days != due_days):
            status = "review"
        else:
            status = "normal"

    rule_trace = {
        "version": "v1",
        "policy_mode": mode,
        "base_rule": base_rule,
        "effective_punch_days": punch_days,
        "due_days": due_days,
        "adjustment_days": adjustment,
        "pending_cross_day_suspicions": pending_suspicions,
        "actual_formula": "base_actual + adjustment_days",
    }

    store.execute(
        "UPDATE attendance_result SET punch_days = ?, due_days = ?, rest_days = ?, actual_days = ?, "
        "status = ?, rule_trace = ?, updated_at = ? WHERE id = ?",
        (float(punch_days), due_days, rest_days, actual_days, status, json_dumps(rule_trace), _now(), result_id),
    )
    return store.query_one("SELECT * FROM attendance_result WHERE id = ?", (result_id,), table="attendance_result")


def recalculate_batch(store: AttendanceStore, batch_id: int) -> None:
    batch = store.query_one("SELECT * FROM import_batch WHERE id = ?", (batch_id,))
    month_end = date(batch["year"], batch["month"], calendar.monthrange(batch["year"], batch["month"])[1])
    active_employees = store.query(
        "SELECT id FROM employee WHERE active = 1 AND (join_date IS NULL OR join_date <= ?)",
        (month_end.isoformat(),),
    )
    for employee in active_employees:
        existing = store.query_one(
            "SELECT id FROM attendance_result WHERE batch_id = ? AND employee_id = ?",
            (batch_id, employee["id"]),
        )
        if existing:
            recalculate_result(store, existing["id"])
        else:
            result_id = store.execute(
                "INSERT INTO attendance_result (batch_id, employee_id, status, updated_at) VALUES (?, ?, 'review', ?)",
                (batch_id, employee["id"], _now()),
            )
            recalculate_result(store, result_id)


def resolve_cross_day(store: AttendanceStore, suspicion_id: int, resolution: str, user_id: int | None) -> dict:
    if resolution not in {"assign_previous", "keep_current"}:
        raise ValueError("无效的审核结果")
    suspicion = store.query_one(
        "SELECT * FROM cross_day_suspicion WHERE id = ?", (suspicion_id,), table="cross_day_suspicion"
    )
    if not suspicion:
        raise ValueError("跨日疑似记录不存在")
    now = _now()
    store.execute(
        "UPDATE cross_day_suspicion SET status = ?, reviewed_by_id = ?, reviewed_at = ? WHERE id = ?",
        (resolution, user_id, now, suspicion_id),
    )
    keep = 1 if resolution == "keep_current" else 0
    store.execute(
        "UPDATE raw_punch_day SET effective_has_punch = ? WHERE id = ?",
        (keep, suspicion["raw_day_id"]),
    )
    if suspicion["employee_id"]:
        result = store.query_one(
            "SELECT id FROM attendance_result WHERE batch_id = ? AND employee_id = ?",
            (suspicion["batch_id"], suspicion["employee_id"]),
        )
        if result:
            recalculate_result(store, result["id"])
    return store.query_one(
        "SELECT * FROM cross_day_suspicion WHERE id = ?", (suspicion_id,), table="cross_day_suspicion"
    )
